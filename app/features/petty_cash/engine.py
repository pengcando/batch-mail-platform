# -*- coding: utf-8 -*-
"""
对账单批量作业 —— 核心业务引擎
================================

设计原则
--------
1. **不驱动本机 Excel**。传统流程靠宏/PAD 操作 Excel，服务端没有 Excel、也不该有。
   这里把原流程里所有需要 Excel 求值的地方，全部显式算成 Python 值：
     - `对账单` 页原本是纯公式页（`='账单数据-原始数据'!X`），这里按列映射直接算值；
     - `账单数据-原始数据!U`（佣金开票金额）与 `!V`（发票开票方）两条业务公式，见
       `_invoice_fee` / `_invoice_party`；
     - `开票信息!C6:C10`（邮件正文用的金额），见 `build_invoice_context`。
   因此产出的 xlsx **不含任何跨表引用**，也就不会出现失效的 #REF!。

2. **只搬"版式"不搬"公式"**。表头文字、字体、列宽、数字格式全部从源模板取快照，
   保证出品的对账单长得和模板一致（详见 core/xlsx_style.py）。

3. **产物纯内存拼装**，最后一次性落 ZIP，符合「服务端没有桌面目录」的约束。

输入文件约定
----------------
M = 商户名单文件（含收件人/抄送/进度列的工作表固定叫 `Sheet1`）
T = 抽数工具文件（含城市明细页、对账单模板页、增减项页、开票信息页）
"""

from __future__ import annotations

import calendar
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook, load_workbook

from app.core import email_html
from app.core.xlsx_style import SheetTemplate, build_sheet, reorder_rows_by_key, set_sheet_number_format

# ---------------------------------------------------------------------------
# 常量：列定义（全部为 1 基序号，与 openpyxl 对齐）
# ---------------------------------------------------------------------------

M_SHEET = "Sheet1"
NCOL_M = 7            # A..G
M_COL_NAME, M_COL_INVOICE_INFO, M_COL_INVOICE_TYPE, M_COL_REMARK, M_COL_TO, M_COL_CC, M_COL_PROGRESS = range(1, 8)

CITY_SHEET_NCOL = 19  # A..S
CITY_COL_CITY, CITY_COL_CUSTOMER, CITY_COL_OM, CITY_COL_SHOP, CITY_COL_BIZ = 1, 2, 3, 4, 5
CITY_COL_SALES, CITY_COL_RATE, CITY_COL_ADJUST, CITY_COL_SETTLE, CITY_COL_REAL_RATE = 6, 7, 8, 9, 10
CITY_COL_DELIVERY, CITY_COL_FOOD, CITY_COL_OTHER, CITY_COL_PAYABLE, CITY_COL_PAYTERM = 11, 12, 13, 14, 15
CITY_COL_CUST_INV, CITY_COL_CUST_CONTRACT, CITY_COL_MCH_INV, CITY_COL_SIGN = 16, 17, 18, 19

STATEMENT_SHEET = "对账单"
STATEMENT_NCOL = 18

ADJUST_SHEET = "增减项"
ADJUST_NCOL = 6       # A..F（城市/情形/通道/商户/计算方法/比例金额）
ADJUST_COL_CITY, ADJUST_COL_SHOP = 1, 4

INVOICE_SHEET = "开票信息"
DATA_SHEET = "账单数据-原始数据"

STATUS_NO_FLOW = "本月无流水"
STATUS_DONE = "已完成"
STATUS_FAILED = "未完成"

# 邮件主题 / 附件名的短横位置不同（`对账单-年月` vs `对账单年月`）是有意设计，不要统一
# 「公司名」请替换成你自己的品牌前缀
SUBJECT_TEMPLATE = "公司名&{merchant}-对账单{ym}"
ATTACH_TEMPLATE = "公司名&{merchant}对账单-{ym}.xlsx"

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


# ---------------------------------------------------------------------------
# 小工具
# ---------------------------------------------------------------------------

def _s(value) -> str:
    return "" if value is None else str(value).strip()


def _num(value):
    """把单元格值尽量转成数字；转不了返回 None。"""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100.0
        except ValueError:
            return None
    try:
        return float(text)
    except ValueError:
        return None


def split_emails(raw) -> list[str]:
    """
    收件人清洗。
    实测 E/F 列可能一个单元格塞多个邮箱、且中英文分号混用（`someone@company.com；`），
    所以按 `; ； , ，` 与空白一起切分，再逐个校验格式、去重。
    """
    parts = re.split(r"[;；,，\s]+", _s(raw))
    out: list[str] = []
    for part in parts:
        part = part.strip().strip(";；,，")
        if part and _EMAIL_RE.match(part) and part not in out:
            out.append(part)
    return out


def brand_of(shop_name: str) -> str:
    """对应 账单数据!B 的公式：`=IF(ISERROR(FIND("（",F)),F,LEFT(F,FIND("（",F)-1))`。"""
    name = _s(shop_name)
    idx = name.find("（")
    return name if idx <= 0 else name[:idx]


def compute_month_value(mode: str, ym: str) -> date:
    """
    `对账单` 的「月份」列。

    原流程该列是公式 `=EDATE(TODAY(),-1)`（运行日期往前一个月），显示格式 `yyyy"年"m"月"`。
    mode = "run_date"（默认，复刻原流程） / "ym"（直接取本次账期）。
    """
    if mode == "ym" and ym and len(ym) >= 7:
        try:
            return date(int(ym[:4]), int(ym[5:7]), 1)
        except ValueError:
            pass
    today = date.today()
    year, month = today.year, today.month - 1
    if month == 0:
        year, month = year - 1, 12
    day = min(today.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


# ---------------------------------------------------------------------------
# 业务公式（原 Excel 里由公式承担的部分）
# ---------------------------------------------------------------------------

def _is_yes(value) -> bool:
    return _s(value) == "是"


def _is_no(value) -> bool:
    return _s(value) == "否"


def _invoice_fee(sales, settle, cust_inv, mch_inv):
    """
    账单数据!U（对账单的「服务费开票金额」）
    公式： =IF(AND(T:T="是", R:R="否"), H-K, 0)
    即：仅当「商户开票情况=是」且「客户开票情况=否」时，取 调整后销售额-调整后结算额，否则记 0。
    """
    if _is_yes(mch_inv) and _is_no(cust_inv):
        sales_num, settle_num = _num(sales), _num(settle)
        if sales_num is not None and settle_num is not None:
            return round(sales_num - settle_num, 6)
        return 0
    return 0


def _invoice_party(cust_inv):
    """账单数据!V（对账单的「客户发票开票方」）：`=IF(R3="否","贵司","公司名")`。

    「公司名」请替换成你自己的品牌名。
    """
    return "贵司" if _is_no(cust_inv) else "公司名"


def build_statement_row(city_row: list, month_value: date) -> list:
    """城市明细行的 19 列 → 对账单的 18 列（逐列对应关系见流程文档 §3.2(2)）。"""
    cust_inv = city_row[CITY_COL_CUST_INV - 1]
    mch_inv = city_row[CITY_COL_MCH_INV - 1]
    return [
        month_value,                              # A 月份 ← EDATE(TODAY(),-1)
        city_row[CITY_COL_CUSTOMER - 1],          # B 客户名称
        city_row[CITY_COL_OM - 1],                # C 营运经理
        city_row[CITY_COL_SHOP - 1],              # D 供餐商户
        city_row[CITY_COL_BIZ - 1],               # E 业务类型
        city_row[CITY_COL_SALES - 1],             # F 调整后销售额
        city_row[CITY_COL_RATE - 1],              # G 合同佣金率
        city_row[CITY_COL_ADJUST - 1],            # H 增减项
        city_row[CITY_COL_SETTLE - 1],            # I 调整后结算额
        city_row[CITY_COL_REAL_RATE - 1],         # J 实际佣金率
        city_row[CITY_COL_DELIVERY - 1],          # K 商户配送费
        city_row[CITY_COL_FOOD - 1],              # L 供应链食材扣款
        city_row[CITY_COL_OTHER - 1],             # M 其他扣减项
        city_row[CITY_COL_PAYABLE - 1],           # N 抵扣后实付金额
        _invoice_fee(city_row[CITY_COL_SALES - 1], city_row[CITY_COL_SETTLE - 1], cust_inv, mch_inv),  # O 服务费开票金额
        city_row[CITY_COL_CUST_CONTRACT - 1],     # P 客户合同情况
        _invoice_party(cust_inv),                 # Q 客户发票开票方
        city_row[CITY_COL_SIGN - 1],              # R 商户签约主体
    ]


def build_invoice_context(merchant_row: list, statement_rows: list, tpl_intro: str, tpl_tip: str) -> dict:
    """
    组装邮件正文所需数据（对应 开票信息!A1:I10）。
    C6 品牌名称 ← 账单数据!B3；C7/C8/C9 ← 开票计算!C7/C8/C9；C10 强制 0.00（不开票）。
    """
    sales_sum = sum(x for x in (_num(r[5]) for r in statement_rows) if x is not None)
    settle_sum = sum(x for x in (_num(r[8]) for r in statement_rows) if x is not None)

    return {
        "intro": tpl_intro,
        "tip": tpl_tip,
        "merchant": merchant_row[M_COL_NAME - 1],
        "sales": round(sales_sum, 2),
        "settle": round(settle_sum, 2),
        "service": round(sales_sum - settle_sum, 2),
        "invoice_fee": 0.0,          # 服务费开票金额置零
        "rows": [
            ("商户品牌名称：", brand_of(merchant_row[M_COL_NAME - 1]), False),
            ("本期销售额：", round(sales_sum, 2), True),
            ("本期结算额：", round(settle_sum, 2), True),
            ("服务费金额：", round(sales_sum - settle_sum, 2), True),
            ("服务费开票金额：", 0.0, True),
        ],
    }


# ---------------------------------------------------------------------------
# 源文件读取
# ---------------------------------------------------------------------------

class SourceError(Exception):
    pass


@dataclass
class Merchant:
    excel_row: int                 # 在 M.Sheet1 中的绝对行号（状态回写用）
    name: str
    to_raw: str
    cc_raw: str
    invoice_info: str = ""
    invoice_type: str = ""
    remark: str = ""

    @property
    def to_list(self) -> list[str]:
        return split_emails(self.to_raw)

    @property
    def cc_list(self) -> list[str]:
        return split_emails(self.cc_raw)


def open_m(path_or_bytes):
    try:
        return load_workbook(path_or_bytes, data_only=False)
    except Exception as exc:                                     # noqa: BLE001
        raise SourceError(f"《不开票商户明细》无法打开：{exc}") from exc


def open_t(path_or_bytes):
    try:
        return load_workbook(path_or_bytes, data_only=False)
    except Exception as exc:                                     # noqa: BLE001
        raise SourceError(f"《商户对账工具》无法打开：{exc}") from exc


def list_city_sheets(wb_t) -> list[dict]:
    """城市工作表 = 名字不是固定功能页的那几个 19 列表页。"""
    reserved = {DATA_SHEET, STATEMENT_SHEET, INVOICE_SHEET, ADJUST_SHEET, "开票计算", "开票计算 "}
    out = []
    for name in wb_t.sheetnames:
        if name in reserved:
            continue
        ws = wb_t[name]
        if ws.max_column < CITY_SHEET_NCOL:
            continue
        out.append({"name": name, "rows": max(ws.max_row - 1, 0)})
    return out


def read_city_rows(wb_t, city: str) -> list[list]:
    if city not in wb_t.sheetnames:
        available = [s["name"] for s in list_city_sheets(wb_t)]
        raise SourceError(f"《商户对账工具》里没有名为「{city}」的工作表。可用城市：{'、'.join(available)}")
    ws = wb_t[city]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=CITY_SHEET_NCOL,
                            values_only=True):
        if all(v is None or _s(v) == "" for v in row):
            continue
        rows.append(list(row))
    return rows


def read_merchants(wb_m, skip_done: bool) -> list[Merchant]:
    if M_SHEET not in wb_m.sheetnames:
        raise SourceError(f"《不开票商户明细》里找不到固定工作表「{M_SHEET}」，实际为：{wb_m.sheetnames}")
    ws = wb_m[M_SHEET]
    merchants: list[Merchant] = []
    for r in range(2, ws.max_row + 1):
        name = _s(ws.cell(row=r, column=M_COL_NAME).value)
        if not name:
            continue
        progress = _s(ws.cell(row=r, column=M_COL_PROGRESS).value)
        if skip_done and progress:
            continue
        merchants.append(
            Merchant(
                excel_row=r,
                name=name,
                to_raw=_s(ws.cell(row=r, column=M_COL_TO).value),
                cc_raw=_s(ws.cell(row=r, column=M_COL_CC).value),
                invoice_info=_s(ws.cell(row=r, column=M_COL_INVOICE_INFO).value),
                invoice_type=_s(ws.cell(row=r, column=M_COL_INVOICE_TYPE).value),
                remark=_s(ws.cell(row=r, column=M_COL_REMARK).value),
            )
        )
    # A11：按 A 列降序
    merchants.sort(key=lambda m: m.name, reverse=True)
    return merchants


# ---------------------------------------------------------------------------
# 输出工作簿
# ---------------------------------------------------------------------------

@dataclass
class Templates:
    statement: SheetTemplate
    adjust: SheetTemplate
    invoice_intro: str = ""
    invoice_tip: str = ""


def load_templates(wb_t) -> Templates:
    if STATEMENT_SHEET not in wb_t.sheetnames:
        raise SourceError(f"《商户对账工具》缺少工作表「{STATEMENT_SHEET}」")
    if ADJUST_SHEET not in wb_t.sheetnames:
        raise SourceError(f"《商户对账工具》缺少工作表「{ADJUST_SHEET}」")

    invoice_intro, invoice_tip = "", ""
    if INVOICE_SHEET in wb_t.sheetnames:
        ws = wb_t[INVOICE_SHEET]
        invoice_intro = ws.cell(row=1, column=1).value or ""
        invoice_tip = ws.cell(row=5, column=1).value or ""

    return Templates(
        statement=SheetTemplate(wb_t[STATEMENT_SHEET], STATEMENT_NCOL),
        adjust=SheetTemplate(wb_t[ADJUST_SHEET], ADJUST_NCOL),
        invoice_intro=invoice_intro,
        invoice_tip=invoice_tip,
    )


def build_statement_workbook(templates: Templates, statement_rows: list[list],
                             adjust_rows: list[list], numeric_format: bool = False) -> bytes:
    """生成一个商户的对账单文件：`对账单` 页 + 追加的 `增减项` 页（流程文档 B10~B23）。"""
    wb = Workbook()
    wb.remove(wb.active)                     # B12：干掉新建工作簿自带的默认 Sheet1

    ws_statement = build_sheet(wb, templates.statement, statement_rows, title=STATEMENT_SHEET)
    if numeric_format:
        set_sheet_number_format(ws_statement, "0.00")

    build_sheet(wb, templates.adjust, adjust_rows, title=ADJUST_SHEET)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_adjust_rows(wb_t, city: str, merchant_name: str) -> tuple[list[list], int]:
    """
    增减项（流程文档 B-3）。
    口径（v6 确认）：m = 源表 `T.增减项` 中「A 列=城市 且 D 列=商户」的匹配行数。
    分支只影响"贴多少行"：m==0 → 仅表头；m>0 → 匹配到的数据行。
    ⚠️ 取区域时按实际有值的列（F 列）判定，绝不能碰 UsedRange —— 实测该表被格式化到 XFD 列。
    """
    ws = wb_t[ADJUST_SHEET]
    matched = []
    for r in range(2, ws.max_row + 1):
        if _s(ws.cell(row=r, column=ADJUST_COL_CITY).value) != city:
            continue
        if _s(ws.cell(row=r, column=ADJUST_COL_SHOP).value) != merchant_name:
            continue
        matched.append([ws.cell(row=r, column=c).value for c in range(1, ADJUST_NCOL + 1)])
    return matched, len(matched)


# ---------------------------------------------------------------------------
# 运行参数与结果
# ---------------------------------------------------------------------------

@dataclass
class RunOptions:
    city: str
    ym: str
    smtp_password: str = ""
    dry_run: bool = True                 # 默认预览：只生成文件与预览，不发信
    skip_done: bool = True               # 默认跳过 M 表 G 列已有状态的行
    month_mode: str = "run_date"         # run_date | ym
    numeric_format: bool = False
    max_merchants: int = 0
    context: dict = field(default_factory=dict)   # 额外透传给结果表的信息


@dataclass
class MerchantResult:
    merchant: str
    to: str
    cc: str
    status: str
    detail: str = ""
    rows: int = 0
    invoice_rows: int = 0
    file_name: str = ""

    def as_dict(self) -> dict:
        return {
            "merchant": self.merchant,
            "to": self.to,
            "cc": self.cc,
            "status": self.status,
            "detail": self.detail,
            "rows": self.rows,
            "adjustRows": self.invoice_rows,
            "fileName": self.file_name,
        }


def zip_result(files: list[tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files:
            zf.writestr(name, data)
    return buf.getvalue()
